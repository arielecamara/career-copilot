"""Single-column DOCX export. No tables, text boxes, or extra columns."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from career_copilot.models import ResumeSentence, SentenceKind
from career_copilot.presentation import PUBLICATION_NOTE, polish
from career_copilot.source_cv import SourceSection

_HEADINGS = {
    SentenceKind.contact: "Contact",
    SentenceKind.summary: "Summary",
    SentenceKind.skill: "Skills",
    SentenceKind.role: "Experience",
    SentenceKind.education: "Education",
    SentenceKind.project: "Projects",
    SentenceKind.publication: "Publications",
    SentenceKind.language: "Languages",
}
_DASH = {SentenceKind.bullet, SentenceKind.skill, SentenceKind.education, SentenceKind.project, SentenceKind.publication}
_COVER_CHROME = ("Hello,", "Sincerely,")


def write_resume(
    path: Path,
    sentences: list[ResumeSentence],
    *,
    header_contact: str | None = None,
    header_name: str | None = None,
    publication_note: bool = False,
) -> None:
    """One column and no tables. Header lines are not evidence sentences."""

    document = _document()
    contact = _optional_line(header_contact)
    if header_name and not any(sentence.kind == SentenceKind.name for sentence in sentences):
        _add_name(document, header_name)
        _add_contact(document, contact)
        contact = None
    last_publication = _last_index(sentences, SentenceKind.publication)
    seen: set[SentenceKind] = set()
    for index, sentence in enumerate(sentences):
        heading = _HEADINGS.get(sentence.kind)
        if heading is not None and sentence.kind not in seen:
            seen.add(sentence.kind)
            document.add_heading(heading, level=1)
        shown = polish(sentence.text)
        if not shown:
            continue
        if sentence.kind == SentenceKind.name:
            _add_name(document, shown)
            if contact:
                _add_contact(document, contact)
                contact = None
            continue
        paragraph = document.add_paragraph()
        if sentence.kind == SentenceKind.role:
            run = paragraph.add_run(shown)
            run.bold = True
        elif sentence.kind in _DASH:
            paragraph.add_run(f"- {shown}")
        else:
            paragraph.add_run(shown)
        if publication_note and index == last_publication:
            document.add_paragraph(PUBLICATION_NOTE)
    _save(document, path)


def write_source_resume(path: Path, sections: tuple[SourceSection, ...] | list[SourceSection]) -> None:
    """One column. Headings and frozen lines are the original CV text."""

    document = _document()
    for section in sections:
        if section.heading:
            document.add_heading(section.heading, level=1)
        for line in section.lines:
            shown = line.text.strip()
            if not shown:
                continue
            if line.style == "name":
                paragraph = document.add_paragraph()
                run = paragraph.add_run(shown)
                run.bold = True
                run.font.size = Pt(16)
                run.font.color.rgb = RGBColor(0, 0, 0)
                continue
            paragraph = document.add_paragraph()
            text = f"- {shown}" if line.bullet else shown
            run = paragraph.add_run(text)
            if line.style == "role":
                run.bold = True
    _save(document, path)


def write_cover_letter(path: Path, sentences: list[ResumeSentence], name: str) -> None:
    document = _document()
    shown_name = polish(name) if isinstance(name, str) else ""
    if shown_name:
        run = document.add_paragraph().add_run(shown_name)
        run.bold = True
    document.add_paragraph(_COVER_CHROME[0])
    for sentence in sentences:
        shown = polish(sentence.text)
        if shown:
            document.add_paragraph(shown)
    document.add_paragraph(_COVER_CHROME[1])
    if shown_name:
        document.add_paragraph(shown_name)
    _save(document, path)


def _document() -> Document:
    document = Document()
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    cols = section._sectPr.find(qn("w:cols"))
    if cols is not None:
        cols.set(qn("w:num"), "1")
    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    heading = document.styles["Heading 1"]
    heading.font.name = "Arial"
    heading.font.color.rgb = RGBColor(0, 0, 0)
    heading.font.size = Pt(14)
    return document


def _add_name(document: Document, text: str) -> None:
    shown = polish(text)
    if not shown:
        return
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = paragraph.add_run(shown)
    run.bold = True
    run.font.size = Pt(16)
    run.font.color.rgb = RGBColor(0, 0, 0)


def _add_contact(document: Document, text: str | None) -> None:
    if text:
        document.add_paragraph(text)


def _optional_line(text: str | None) -> str | None:
    if not isinstance(text, str):
        return None
    shown = polish(text)
    return shown or None


def _last_index(sentences: list[ResumeSentence], kind: SentenceKind) -> int | None:
    found: int | None = None
    for index, sentence in enumerate(sentences):
        if sentence.kind == kind:
            found = index
    return found


def _save(document: Document, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(path)

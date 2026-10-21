from pathlib import Path

from docx import Document

from career_copilot.cli import main
from career_copilot.identity import decide_match
from career_copilot.models import ConfirmationState, EvidenceCategory, ExactDate
from career_copilot.profile import load_profile, save_profile
from tests.factories import atom, document, profile_for


def _clear(**overrides: object):
    payload: dict[str, object] = {
        "confirmation": ConfirmationState.confirmed,
        "category": EvidenceCategory.mentioned_only,
        "employer": None,
        "title": None,
        "start_date": None,
        "end_date": None,
    }
    payload.update(overrides)
    return atom(**payload)


def _profile(profile_id: str, atoms: list):
    record = document([("paragraph", atoms[0].statement)])
    return profile_for(record, atoms).model_copy(update={"profile_id": profile_id})


def test_email_phone_and_employment_are_strong_and_a_name_is_not():
    email = _clear(evidence_id="ev_" + "a" * 16, statement="ada@example.com", source_excerpt="ada@example.com")
    phone = _clear(
        evidence_id="ev_" + "b" * 16,
        statement="Phone +1 555-123-4567",
        source_excerpt="Phone +1 555-123-4567",
    )
    named = _clear(evidence_id="ev_" + "c" * 16, statement="Ada Lovelace", source_excerpt="Ada Lovelace")
    role = atom(
        evidence_id="ev_" + "d" * 16,
        confirmation=ConfirmationState.confirmed,
        statement="Engineer | Acme 2020 – 2022",
        source_excerpt="Engineer | Acme 2020 – 2022",
        employer="Acme",
        title="Engineer",
        start_date=ExactDate(raw="2020", year=2020),
        end_date=ExactDate(raw="2022", year=2022),
    )
    by_email = _profile("pf_" + "1" * 16, [email])
    by_phone = _profile("pf_" + "2" * 16, [phone])
    by_name = _profile("pf_" + "3" * 16, [named])
    by_role = _profile("pf_" + "4" * 16, [role])

    assert decide_match([email], [by_email, by_name]).action == "attach"
    assert decide_match([email], [by_email, by_name]).reason == "email"
    assert decide_match([phone], [by_phone]).reason == "phone"
    assert decide_match([named], [by_name]).action == "ask"
    assert decide_match([named], [by_name]).reason == "only a name matches"
    assert decide_match([role], [by_role]).reason == "overlapping employment"
    other = atom(
        evidence_id="ev_" + "e" * 16,
        confirmation=ConfirmationState.confirmed,
        statement="Engineer | OtherCo 2020 – 2022",
        source_excerpt="Engineer | OtherCo 2020 – 2022",
        employer="OtherCo",
        title="Engineer",
        start_date=ExactDate(raw="2020", year=2020),
        end_date=ExactDate(raw="2022", year=2022),
    )
    assert decide_match([other], [by_role]).action == "create"
    assert decide_match([email], [by_email, _profile("pf_" + "5" * 16, [email])]).action == "ask"

    years = _clear(evidence_id="ev_" + "f" * 16, statement="2019-2020", source_excerpt="2019-2020")
    assert decide_match([years], [_profile("pf_" + "6" * 16, [years])]).action == "create"


def _cv(path: Path, email: str, employer: str) -> None:
    document = Document()
    document.add_paragraph("Ada Lovelace")
    document.add_paragraph(email)
    document.add_heading("Professional Experience", level=1)
    document.add_paragraph(f"Engineer | {employer} 2020 – 2022")
    document.save(path)


def test_ingest_attaches_or_creates_a_separate_profile(tmp_path: Path, capsys):
    root = tmp_path / "profiles"
    first_id = "pf_" + "a" * 16
    email = _clear(evidence_id="ev_" + "a" * 16, statement="ada@example.com", source_excerpt="ada@example.com")
    save_profile(root / first_id / "profile.json", _profile(first_id, [email]))
    before = (root / first_id / "profile.json").read_bytes()

    named = tmp_path / "named.docx"
    _cv(named, "other@example.com", "OtherCo")
    created = main(["ingest", str(named), "--profiles-dir", str(root), "--choose", "new"])
    assert created == 0
    assert "Created profile" in capsys.readouterr().out
    assert (root / first_id / "profile.json").read_bytes() == before
    created_dirs = [path for path in root.iterdir() if path.is_dir() and path.name != first_id]
    assert len(created_dirs) == 1
    assert load_profile(created_dirs[0] / "profile.json").profile_id == created_dirs[0].name

    matched = tmp_path / "matched.docx"
    _cv(matched, "ada@example.com", "FarCo")
    attached = main(["ingest", str(matched), "--profiles-dir", str(root)])
    assert attached == 0
    output = capsys.readouterr().out
    assert f"Attached to {first_id} (email)." in output
    stored = load_profile(root / first_id / "profile.json")
    assert any(atom.statement == "ada@example.com" for atom in stored.atoms)
    assert len(stored.documents) == 2
    assert load_profile(created_dirs[0] / "profile.json").profile_id != first_id
    assert all("ada@example.com" not in atom.statement for atom in load_profile(created_dirs[0] / "profile.json").atoms)

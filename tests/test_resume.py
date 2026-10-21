from pathlib import Path

import pytest
from docx import Document

from career_copilot.cli import main
from career_copilot.docx_export import write_cover_letter, write_resume
from career_copilot.job import analyze_job
from career_copilot.models import (
    ChecklistTally,
    ConfirmationState,
    ConstraintKind,
    ConstraintResult,
    EvidenceCategory,
    ExactDate,
    ExperienceYears,
    HardConstraint,
    JobAnalysis,
    Residence,
    ResumeSentence,
    SentenceKind,
)
from career_copilot.profile import save_profile
from career_copilot.presentation import (
    COVER_LEAD,
    PUBLICATION_NOTE,
    fact_groups,
    fit_feedback,
    header_contact,
    polish,
)
from career_copilot.resume import build_resume
from career_copilot.validator import validate_resume
from tests.factories import atom, document, profile_for

AS_OF = __import__("datetime").date(2026, 10, 13)


def _confirmed(**overrides: object):
    payload: dict[str, object] = {"confirmation": ConfirmationState.confirmed}
    payload.update(overrides)
    return atom(**payload)


def _clear_role_fields(**overrides: object):
    payload: dict[str, object] = {
        "employer": None,
        "title": None,
        "start_date": None,
        "end_date": None,
        "category": EvidenceCategory.mentioned_only,
    }
    payload.update(overrides)
    return _confirmed(**payload)


def _profile_atoms():
    name = _clear_role_fields(
        evidence_id="ev_" + "1" * 16,
        statement="Ada Lovelace",
        source_excerpt="Ada Lovelace",
    )
    skill = _clear_role_fields(
        evidence_id="ev_" + "2" * 16,
        statement="Python",
        source_excerpt="Python",
        skills=["Python"],
    )
    education = _clear_role_fields(
        evidence_id="ev_" + "3" * 16,
        statement="BSc in History",
        source_excerpt="BSc in History",
        category=EvidenceCategory.coursework,
    )
    older = _confirmed(
        evidence_id="ev_" + "4" * 16,
        statement="Engineer | OldCo 2010 – 2011",
        source_excerpt="Engineer | OldCo 2010 – 2011",
        employer="OldCo",
        title="Engineer",
        start_date=ExactDate(raw="2010", year=2010),
        end_date=ExactDate(raw="2011", year=2011),
    )
    matched = _confirmed(
        evidence_id="ev_" + "5" * 16,
        statement="Used Python to reduce latency by 20%.",
        source_excerpt="Used Python to reduce latency by 20%.",
        employer="OldCo",
        title="Engineer",
        start_date=ExactDate(raw="2010", year=2010),
        end_date=ExactDate(raw="2011", year=2011),
        skills=["Python"],
    )
    unrelated = _confirmed(
        evidence_id="ev_" + "6" * 16,
        statement="Organized the archive.",
        source_excerpt="Organized the archive.",
        employer="OldCo",
        title="Engineer",
        start_date=ExactDate(raw="2010", year=2010),
        end_date=ExactDate(raw="2011", year=2011),
    )
    newer = _confirmed(
        evidence_id="ev_" + "7" * 16,
        statement="Engineer | NewCo 2020 – 2022",
        source_excerpt="Engineer | NewCo 2020 – 2022",
        employer="NewCo",
        title="Engineer",
        start_date=ExactDate(raw="2020", year=2020),
        end_date=ExactDate(raw="2022", year=2022),
    )
    hidden = atom(
        evidence_id="ev_" + "8" * 16,
        statement="Kubernetes",
        source_excerpt="Kubernetes",
        category=EvidenceCategory.mentioned_only,
        employer=None,
        title=None,
        start_date=None,
        end_date=None,
        skills=["Kubernetes"],
    )
    return [name, skill, education, older, matched, unrelated, newer, hidden]


def _analysis(atoms, job: str):
    return analyze_job(job, atoms, as_of=AS_OF)


def test_resume_keeps_profile_order_and_refuses_a_missing_skill():
    atoms = _profile_atoms()
    job = "Requirements:\n- Python\n- Kubernetes"
    analysis = _analysis(atoms, job)
    bundle = build_resume(atoms, analysis, cover_letter=True)
    texts = [sentence.text for sentence in bundle.sentences]
    assert texts.index("Engineer | OldCo 2010 – 2011") < texts.index("Engineer | NewCo 2020 – 2022")
    assert texts.index("Engineer | NewCo 2020 – 2022") < texts.index("BSc in History")
    assert "Used Python to reduce latency by 20%." in texts
    assert "Organized the archive." not in texts
    assert "Kubernetes" not in "\n".join(texts)
    assert "Kubernetes" not in "\n".join(sentence.text for sentence in bundle.cover)
    assert "Kubernetes" in bundle.diff_markdown
    assert "Omitted" in bundle.diff_markdown
    profile = profile_for(document([("paragraph", "Python")]), atoms)
    assert validate_resume([*bundle.sentences, *bundle.cover], profile, analysis) == []


def test_adversarial_missing_skill_is_absent_from_the_resume_and_cover_letter():
    atoms = _profile_atoms()
    docker = _clear_role_fields(
        evidence_id="ev_" + "9" * 16,
        statement="Docker",
        source_excerpt="Docker",
        skills=["Docker"],
    )
    atoms = [*atoms, docker]
    job = "Requirements:\n- Kubernetes\n- Docker"
    analysis = _analysis(atoms, job)
    bundle = build_resume(atoms, analysis, cover_letter=True)
    visible = "\n".join(sentence.text for sentence in [*bundle.sentences, *bundle.cover])
    assert "docker" in visible.casefold()
    assert "kubernetes" not in visible.casefold()
    profile = profile_for(document([("paragraph", "Docker")]), atoms)
    assert validate_resume([*bundle.sentences, *bundle.cover], profile, analysis) == []
    poisoned = ResumeSentence(
        text="Used Kubernetes to ship the service.",
        evidence_ids=["ev_" + "9" * 16],
        kind=SentenceKind.bullet,
    )
    violations = validate_resume([poisoned], profile, analysis)
    assert any(item.code == "absent_skill" for item in violations)


def test_validator_rejects_changed_facts_and_a_stronger_verb():
    atoms = _profile_atoms()
    profile = profile_for(document([("paragraph", "Used Python to reduce latency by 20%.")]), atoms)
    analysis = _analysis(atoms, "Requirements:\n- Python")
    number = ResumeSentence(
        text="Used Python to reduce latency by 50%.",
        evidence_ids=["ev_" + "5" * 16],
        kind=SentenceKind.bullet,
    )
    verb = ResumeSentence(
        text="Led the migration and used Python to reduce latency by 20%.",
        evidence_ids=["ev_" + "5" * 16],
        kind=SentenceKind.bullet,
    )
    employer = ResumeSentence(
        text="Engineer | Acme Corp | 2010 – 2011",
        evidence_ids=["ev_" + "4" * 16],
        kind=SentenceKind.role,
    )
    assert any(item.code == "number" for item in validate_resume([number], profile, analysis))
    assert any(item.code == "ownership" for item in validate_resume([verb], profile, analysis))
    assert any(item.code == "employer" for item in validate_resume([employer], profile, analysis))


def test_postgres_may_be_rewritten_as_postgresql_and_docker_may_not():
    postgres = _clear_role_fields(
        evidence_id="ev_" + "a" * 16,
        statement="Postgres",
        source_excerpt="Postgres",
        skills=["Postgres"],
    )
    analysis = _analysis([postgres], "Requirements:\n- PostgreSQL")
    bundle = build_resume([postgres], analysis)
    assert bundle.sentences[0].text == "PostgreSQL"
    assert bundle.sentences[0].evidence_ids == ["ev_" + "a" * 16]
    assert "PostgreSQL" in bundle.diff_markdown
    assert "Original skill order: Postgres." in bundle.diff_markdown
    profile = profile_for(document([("paragraph", "Postgres")]), [postgres])
    assert validate_resume(bundle.sentences, profile, analysis) == []

    bullet = _confirmed(
        evidence_id="ev_" + "c" * 16,
        statement="Used Postgres to store records.",
        source_excerpt="Used Postgres to store records.",
        skills=["Postgres"],
    )
    rewritten = build_resume([bullet], _analysis([bullet], "Requirements:\n- PostgreSQL"))
    assert rewritten.sentences[0].text == "Used PostgreSQL to store records."
    assert "Used PostgreSQL to store records." in rewritten.diff_markdown
    assert "Used Postgres to store records." in rewritten.diff_markdown
    bullet_profile = profile_for(document([("paragraph", "Used Postgres to store records.")]), [bullet])
    assert validate_resume(rewritten.sentences, bullet_profile, _analysis([bullet], "Requirements:\n- PostgreSQL")) == []

    docker = _clear_role_fields(
        evidence_id="ev_" + "b" * 16,
        statement="Docker",
        source_excerpt="Docker",
        skills=["Docker"],
    )
    refused = build_resume([docker], _analysis([docker], "Requirements:\n- Kubernetes"))
    assert all(sentence.text != "Kubernetes" for sentence in refused.sentences)
    assert "Kubernetes" in refused.diff_markdown


def test_docx_is_single_column_and_the_cli_writes_the_diff(tmp_path: Path):
    atoms = _profile_atoms()
    profile = profile_for(document([("paragraph", "Ada Lovelace")]), atoms)
    analysis = _analysis(atoms, "Requirements:\n- Python\n- Kubernetes")
    bundle = build_resume(atoms, analysis, cover_letter=True)
    resume_path = tmp_path / "resume.docx"
    cover_path = tmp_path / "cover.docx"
    write_resume(resume_path, bundle.sentences)
    write_cover_letter(cover_path, bundle.cover, "Ada Lovelace")
    resume = Document(resume_path)
    cover = Document(cover_path)
    resume_text = "\n".join(paragraph.text for paragraph in resume.paragraphs)
    cover_text = "\n".join(paragraph.text for paragraph in cover.paragraphs)
    assert resume.tables == []
    assert cover.tables == []
    assert 'w:num="2"' not in resume.sections[0]._sectPr.xml
    assert "OldCo" in resume_text and resume_text.index("OldCo") < resume_text.index("NewCo")
    assert "Kubernetes" not in resume_text
    assert "Kubernetes" not in cover_text
    assert "Used Python to reduce latency by 20%." in cover_text

    profile_path = tmp_path / "profile.json"
    save_profile(profile_path, profile)
    before = profile_path.read_bytes()
    out = tmp_path / "out.docx"
    diff = tmp_path / "out.diff.md"
    code = main(
        [
            "resume",
            "--text",
            "Requirements:\n- Python\n- Kubernetes",
            "--profile",
            str(profile_path),
            "--as-of",
            "2026-10-13",
            "--out",
            str(out),
            "--diff",
            str(diff),
            "--cover-letter",
            str(tmp_path / "letter.docx"),
            "--confirm",
            "resume,cover",
        ]
    )
    assert code == 0
    assert profile_path.read_bytes() == before
    assert out.is_file() and diff.is_file()
    assert "Kubernetes" in diff.read_text(encoding="utf-8")
    assert "Kubernetes" not in "\n".join(paragraph.text for paragraph in Document(out).paragraphs)


def test_cover_letter_uses_one_confirmed_sentence_per_matched_requirement():
    atoms = _profile_atoms()
    graph = _confirmed(
        evidence_id="ev_" + "d" * 16,
        statement="Used Neo4j to store the graph.",
        source_excerpt="Used Neo4j to store the graph.",
        skills=["Neo4j"],
        employer="OldCo",
        title="Engineer",
        start_date=ExactDate(raw="2010", year=2010),
        end_date=ExactDate(raw="2011", year=2011),
    )
    atoms = [*atoms, graph]
    bundle = build_resume(atoms, _analysis(atoms, "Requirements:\n- Python\n- Neo4j"), cover_letter=True)
    cover = [sentence.text for sentence in bundle.cover]
    assert cover == [
        f"{COVER_LEAD}Used Python to reduce latency by 20%.",
        f"{COVER_LEAD}Used Neo4j to store the graph.",
    ]
    assert all("Kubernetes" not in sentence.text for sentence in bundle.cover)
    profile = profile_for(document([("paragraph", "Ada Lovelace")]), atoms)
    assert validate_resume(bundle.cover, profile, _analysis(atoms, "Requirements:\n- Python\n- Neo4j")) == []


def test_resume_writes_only_what_the_user_confirms(tmp_path: Path, monkeypatch: pytest.MonkeyPatch):
    atoms = _profile_atoms()
    profile = profile_for(document([("paragraph", "Ada Lovelace")]), atoms)
    profile_path = tmp_path / "profile.json"
    save_profile(profile_path, profile)
    out = tmp_path / "out.docx"
    cover = tmp_path / "letter.docx"
    answers = iter(["n", "n"])
    monkeypatch.setattr("builtins.input", lambda _prompt: next(answers))
    declined = main(
        [
            "resume",
            "--text",
            "Requirements:\n- Python",
            "--profile",
            str(profile_path),
            "--as-of",
            "2026-10-13",
            "--out",
            str(out),
            "--cover-letter",
            str(cover),
        ]
    )
    assert declined == 0
    assert not out.exists()
    assert not cover.exists()

    code = main(
        [
            "resume",
            "--text",
            "Requirements:\n- Python",
            "--profile",
            str(profile_path),
            "--as-of",
            "2026-10-13",
            "--out",
            str(out),
            "--cover-letter",
            str(cover),
            "--confirm",
            "resume",
        ]
    )
    assert code == 0
    assert out.is_file()
    assert not cover.exists()
    assert (tmp_path / "out.diff.md").is_file()


def test_generate_is_still_not_a_command():
    with pytest.raises(SystemExit) as caught:
        main(["generate"])
    assert caught.value.code == 2


def _project(evidence_id: str, statement: str):
    return _confirmed(
        evidence_id=evidence_id,
        statement=statement,
        source_excerpt=statement,
        category=EvidenceCategory.project,
        employer=None,
        title=None,
        start_date=None,
        end_date=None,
    )


def test_academic_order_puts_education_before_experience():
    atoms = _profile_atoms()
    analysis = _analysis(atoms, "Requirements:\n- Python")
    academic = [sentence.text for sentence in build_resume(atoms, analysis, academic=True).sentences]
    industry = [sentence.text for sentence in build_resume(atoms, analysis).sentences]
    assert academic.index("BSc in History") < academic.index("Engineer | OldCo 2010 – 2011")
    assert industry.index("Engineer | NewCo 2020 – 2022") < industry.index("BSc in History")


def test_publications_keep_the_two_most_relevant():
    atoms = [
        _project("ev_" + "1" * 16, "Published a paper on archives."),
        _project("ev_" + "2" * 16, "Published a third paper."),
        _project("ev_" + "3" * 16, "Published a Python paper."),
    ]
    bundle = build_resume(atoms, _analysis(atoms, "Requirements:\n- Python"))
    shown = [sentence.text for sentence in bundle.sentences if sentence.kind == SentenceKind.publication]
    assert shown[0] == "Published a Python paper."
    assert len(shown) == 2
    assert "Published a third paper." not in shown
    assert bundle.publication_note is True
    assert PUBLICATION_NOTE not in "\n".join(sentence.text for sentence in bundle.sentences)
    assert "Publications are limited to the two most relevant." in bundle.diff_markdown


def test_cover_letter_stops_at_the_stated_word_limit():
    atoms = _profile_atoms()
    atoms.append(
        _confirmed(
            evidence_id="ev_" + "d" * 16,
            statement="Used Neo4j to store the graph.",
            source_excerpt="Used Neo4j to store the graph.",
            skills=["Neo4j"],
            employer="OldCo",
            title="Engineer",
            start_date=ExactDate(raw="2010", year=2010),
            end_date=ExactDate(raw="2011", year=2011),
        )
    )
    job = "Requirements:\n- Python\n- Neo4j\nCover letter must be 12 words."
    limited = build_resume(atoms, _analysis(atoms, job), cover_letter=True, job_text=job)
    assert len(limited.cover) == 1
    assert limited.cover[0].text.startswith(COVER_LEAD)
    full = build_resume(atoms, _analysis(atoms, "Requirements:\n- Python\n- Neo4j"), cover_letter=True)
    assert len(full.cover) == 2


def test_docx_polishes_text_adds_the_publication_note_and_stays_single_column(tmp_path: Path):
    sentences = [
        ResumeSentence(text="Lisboa aˆ", evidence_ids=["ev_" + "1" * 16], kind=SentenceKind.contact),
        ResumeSentence(text="Published a paper.", evidence_ids=["ev_" + "2" * 16], kind=SentenceKind.publication),
    ]
    path = tmp_path / "resume.docx"
    write_resume(
        path,
        sentences,
        header_contact="City , Portugal",
        header_name="Ada Lovelace",
        publication_note=True,
    )
    document = Document(path)
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert document.tables == []
    assert 'w:num="2"' not in document.sections[0]._sectPr.xml
    assert text.index("Ada Lovelace") < text.index("City, Portugal")
    assert "Lisboa â" in text
    assert polish("Lisboa aˆ") == "Lisboa â"
    assert "Publications" in text
    assert text.index("Published a paper.") < text.index(PUBLICATION_NOTE)
    assert PUBLICATION_NOTE not in sentences[1].text


def test_residence_is_header_text_and_not_a_resume_sentence():
    atoms = _profile_atoms()
    profile = profile_for(document([("paragraph", "Ada Lovelace")]), atoms)
    profile = profile.model_copy(
        update={
            "display_name": "Ada Lovelace",
            "residence": Residence(city="Lisbon", country="Portugal", contact="ada@example.com"),
        }
    )
    assert header_contact(profile) == "ada@example.com | Lisbon | Portugal"
    bundle = build_resume(atoms, _analysis(atoms, "Requirements:\n- Python"))
    visible = "\n".join(sentence.text for sentence in bundle.sentences)
    assert "Lisbon" not in visible
    assert "Portugal" not in visible


def test_polish_repairs_split_accents_without_changing_numbers():
    assert polish("aˆrvore") == "árvore"
    assert polish("Comunicac¸˜ao") == "Comunicação"
    assert polish("Ciˆencia") == "Ciência"
    assert polish("Jo˜ao") == "João"
    assert polish("Funda¸ca˜o") == "Fundação"
    assert polish("Associa¸c˜ao") == "Associação"
    assert polish("Marc¸al") == "Marçal"
    assert polish("Ca¸cador") == "Caçador"
    assert polish("s´ıtio arqueolo´gico") == "sítio arqueológico"
    assert polish("Arque´ologos") == "Arqueólogos"
    assert polish("Sera´ este") == "Será este"
    assert polish("Image • Recognition") == "Image Recognition"
    assert polish("di!erent") == "different"
    assert polish("Hello!") == "Hello!"
    assert polish("2019-2020") == "2019-2020"
    assert polish("latency by 20%") == "latency by 20%"


def test_fact_groups_leave_out_other_facts():
    paper = _project("ev_" + "a" * 16, "Published a paper.")
    name = _clear_role_fields(
        evidence_id="ev_" + "b" * 16,
        statement="Ada Lovelace",
        source_excerpt="Ada Lovelace",
    )
    labels = [label for label, _rows in fact_groups([name, paper])]
    assert labels == ["Publications"]


def test_fit_feedback_uses_the_checklist_and_not_a_score():
    assert fit_feedback(_feedback(0, 0, 0)) == ("Negative", "Not enough information")
    assert fit_feedback(_feedback(1, 0, 2)) == ("Negative", "Missing key skills")
    assert fit_feedback(_feedback(1, 1, 0)) == ("Positive", "Mostly qualified")
    assert fit_feedback(_feedback(2, 0, 0)) == ("Positive", "Well qualified")
    below = _feedback(
        1,
        0,
        0,
        [
            HardConstraint(
                text="10 years",
                kind=ConstraintKind.minimum_years,
                result=ConstraintResult.failed,
                required_years=10,
                observed_days=400,
            )
        ],
    )
    assert fit_feedback(below) == ("Negative", "Below the stated experience")
    blocked = _feedback(
        1,
        0,
        0,
        [HardConstraint(text="Citizen", kind=ConstraintKind.citizenship, result=ConstraintResult.failed)],
    )
    assert fit_feedback(blocked) == ("Negative", "Does not meet a stated requirement")
    over = _feedback(
        1,
        0,
        0,
        [
            HardConstraint(
                text="3 years",
                kind=ConstraintKind.minimum_years,
                result=ConstraintResult.passed,
                required_years=3,
                observed_days=2556,
            )
        ],
    )
    assert fit_feedback(over) == ("Negative", "Overqualified")
    close = _feedback(
        1,
        0,
        0,
        [
            HardConstraint(
                text="5 years",
                kind=ConstraintKind.minimum_years,
                result=ConstraintResult.passed,
                required_years=5,
                observed_days=2556,
            )
        ],
    )
    assert fit_feedback(close) == ("Positive", "Well qualified")


def _feedback(met: int, partial: int, missing: int, constraints: list[HardConstraint] | None = None) -> JobAnalysis:
    judged = met + partial + missing
    percent = None if judged == 0 else round((met + 0.5 * partial) / judged * 100, 1)
    return JobAnalysis(
        years=ExperienceYears(days=400, as_of="2026-10-13", complete=True),
        tally=ChecklistTally(met=met, partial=partial, missing=missing, unknown=0, percent=percent),
        hard_constraints=constraints or [],
    )


def test_docx_keeps_two_publications_and_repairs_their_accents(tmp_path: Path):
    atoms = [
        _project("ev_" + "1" * 16, "aˆrvore published a paper on archives."),
        _project("ev_" + "2" * 16, "Published Comunicac¸˜ao in a journal."),
        _project("ev_" + "3" * 16, "Published a third paper on pottery."),
    ]
    bundle = build_resume(atoms, _analysis(atoms, "Requirements:\n- Python"))
    shown = [sentence for sentence in bundle.sentences if sentence.kind == SentenceKind.publication]
    assert len(shown) == 2
    assert "aˆrvore" in shown[0].text
    path = tmp_path / "resume.docx"
    write_resume(path, bundle.sentences, publication_note=bundle.publication_note)
    text = "\n".join(paragraph.text for paragraph in Document(path).paragraphs)
    assert "árvore" in text
    assert "Comunicação" in text
    assert "aˆrvore" not in text
    assert "third paper" not in text
    assert PUBLICATION_NOTE in text
    assert all(PUBLICATION_NOTE not in sentence.text for sentence in bundle.sentences)

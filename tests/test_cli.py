from pathlib import Path

import pytest
from docx import Document

from career_copilot.cli import main, review_profile
from career_copilot.models import ConfirmationState, DocumentKind
from career_copilot.profile import load_profile


def _write_docx(path: Path, employer: str, end_year: str) -> None:
    document = Document()
    document.add_heading("Professional Experience", level=1)
    document.add_paragraph(f"Engineer | {employer} 2020 – {end_year}")
    document.add_paragraph("Reduced latency by 20%.", style="List Bullet")
    document.add_heading("Technical Skills", level=1)
    document.add_paragraph("Python, Postgres", style="List Bullet")
    document.save(path)


def test_generate_is_not_a_command():
    with pytest.raises(SystemExit) as caught:
        main(["generate"])
    assert caught.value.code == 2


def test_ingest_confirm_edit_and_conflict_gate(tmp_path: Path):
    profile_path = tmp_path / "profile.json"
    first = tmp_path / "one.docx"
    second = tmp_path / "two.docx"
    _write_docx(first, "Acme", "2022")
    _write_docx(second, "Acme", "2023")
    common = ["--profile", str(profile_path)]

    assert main(["ingest", str(first), *common]) == 0
    stored = load_profile(profile_path)
    assert stored.documents[0].filename == "one.docx"
    assert all(atom.confirmation == ConfirmationState.extracted for atom in stored.atoms)
    role = next(atom for atom in stored.atoms if atom.employer == "Acme")
    skill = next(atom for atom in stored.atoms if atom.statement == "Python")

    assert main(["confirm", role.evidence_id, *common]) == 0
    confirmed = load_profile(profile_path).atom(role.evidence_id)
    assert confirmed is not None
    assert confirmed.confirmation == ConfirmationState.confirmed

    assert main(["edit", skill.evidence_id, "--statement", "Python", *common]) == 0
    edited = load_profile(profile_path).atom(skill.evidence_id)
    assert edited is not None
    assert edited.evidence_id == skill.evidence_id
    assert edited.confirmation == ConfirmationState.extracted

    assert main(["ingest", str(second), "--kind", DocumentKind.cv.value, *common]) == 0
    conflicted = load_profile(profile_path)
    assert len(conflicted.documents) == 2
    open_conflicts = [conflict for conflict in conflicted.conflicts if not conflict.resolved]
    assert open_conflicts
    blocked = main(["confirm-all", *common])
    assert blocked == 2
    still = load_profile(profile_path)
    assert still.atom(role.evidence_id).confirmation == ConfirmationState.confirmed
    assert any(atom.confirmation == ConfirmationState.extracted for atom in still.atoms)

    conflict_id = open_conflicts[0].conflict_id
    assert main(["resolve", conflict_id, "--note", "Keep both stints.", *common]) == 0
    resolved = load_profile(profile_path)
    assert all(atom.evidence_id for atom in resolved.atoms)
    assert len(resolved.atoms) == len(conflicted.atoms)
    assert main(["confirm-all", *common]) == 0
    final = load_profile(profile_path)
    assert all(atom.confirmation == ConfirmationState.confirmed for atom in final.atoms)
    assert any(conflict.resolved and conflict.resolution_note == "Keep both stints." for conflict in final.conflicts)
    assert main(["check", *common]) == 0


def test_edit_rejects_a_skill_that_is_not_in_the_excerpt(tmp_path: Path):
    profile_path = tmp_path / "profile.json"
    source = tmp_path / "cv.docx"
    _write_docx(source, "Acme", "2022")
    assert main(["ingest", str(source), "--profile", str(profile_path)]) == 0
    bullet = next(atom for atom in load_profile(profile_path).atoms if "20%" in atom.statement)
    exit_code = main(
        ["edit", bullet.evidence_id, "--skill", "Kubernetes", "--profile", str(profile_path)]
    )
    assert exit_code == 1
    again = load_profile(profile_path).atom(bullet.evidence_id)
    assert again is not None
    assert "Kubernetes" not in again.skills


def test_reingest_keeps_a_resolution_for_the_same_bytes(tmp_path: Path):
    profile_path = tmp_path / "profile.json"
    first = tmp_path / "one.docx"
    second = tmp_path / "two.docx"
    _write_docx(first, "Acme", "2022")
    _write_docx(second, "Acme", "2023")
    common = ["--profile", str(profile_path)]
    assert main(["ingest", str(first), *common]) == 0
    assert main(["ingest", str(second), *common]) == 0
    conflict_id = next(conflict.conflict_id for conflict in load_profile(profile_path).conflicts)
    assert main(["resolve", conflict_id, "--note", "Reviewed.", *common]) == 0
    assert main(["ingest", str(first), *common]) == 0
    conflicts = load_profile(profile_path).conflicts
    assert any(conflict.conflict_id == conflict_id and conflict.resolved for conflict in conflicts)


def test_multi_column_ingest_does_not_write_a_profile(tmp_path: Path, capsys: pytest.CaptureFixture[str]):
    profile_path = tmp_path / "profile.json"
    source = tmp_path / "columns.docx"
    document = Document()
    document.add_paragraph("Two columns")
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    cols = OxmlElement("w:cols")
    cols.set(qn("w:num"), "2")
    document.sections[0]._sectPr.append(cols)
    document.save(source)
    assert main(["ingest", str(source), "--profile", str(profile_path)]) == 1
    assert not profile_path.exists()
    assert "single-column" in capsys.readouterr().err


def test_review_confirms_and_can_edit(tmp_path: Path):
    profile_path = tmp_path / "profile.json"
    source = tmp_path / "cv.docx"
    _write_docx(source, "Acme", "2022")
    assert main(["ingest", str(source), "--profile", str(profile_path)]) == 0
    profile = load_profile(profile_path)
    answers = iter(["c", "e", "Shortened the claim.", "q"])

    def fake_input(_prompt: str) -> str:
        return next(answers)

    updated = review_profile(profile, input_fn=fake_input, write=lambda *_args: None)
    confirmed = [atom for atom in updated.atoms if atom.confirmation == ConfirmationState.confirmed]
    assert len(confirmed) == 1
    edited = next(atom for atom in updated.atoms if atom.statement == "Shortened the claim.")
    assert edited.confirmation == ConfirmationState.extracted
    assert edited.evidence_id == updated.atoms[1].evidence_id

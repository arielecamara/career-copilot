from pathlib import Path

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

from career_copilot.errors import ParseError, UnsupportedLayoutError
from career_copilot.extract import extract_atoms
from career_copilot.grounding import grounding_problems
from career_copilot.models import DocumentKind
from career_copilot.parsers import parse_document
from tests.factories import profile_for


def _write_docx(path: Path, *, columns: int = 1) -> None:
    document = Document()
    document.add_paragraph("Ada Lovelace")
    document.add_heading("Professional Experience", level=1)
    document.add_paragraph("Engineer | Acme 2020 – 2022")
    document.add_paragraph("Led the API migration and cut latency by 20% using Postgres.", style="List Bullet")
    document.add_heading("Technical Skills", level=1)
    document.add_paragraph("Python, Postgres, Docker", style="List Bullet")
    if columns > 1:
        sect_pr = document.sections[0]._sectPr
        cols = sect_pr.find(qn("w:cols"))
        if cols is None:
            cols = OxmlElement("w:cols")
            sect_pr.append(cols)
        cols.set(qn("w:num"), str(columns))
    document.save(path)


def _write_pdf(path: Path, lines: list[str]) -> None:
    pdf = canvas.Canvas(str(path), pagesize=A4)
    _height = A4[1]
    y = _height - 72
    for line in lines:
        pdf.drawString(72, y, line)
        y -= 16
    pdf.save()


def _write_two_column_pdf(path: Path) -> None:
    pdf = canvas.Canvas(str(path), pagesize=A4)
    height = A4[1]
    for line in range(12):
        y = height - 72 - line * 16
        pdf.drawString(50, y, f"Left column sentence {line} alpha beta")
        pdf.drawString(340, y, f"Right column sentence {line} gamma delta")
    pdf.save()


def test_docx_parses_into_grounded_atoms_with_stable_ids(tmp_path: Path):
    path = tmp_path / "cv.docx"
    _write_docx(path)
    first = parse_document(path, DocumentKind.cv)
    second = parse_document(path, DocumentKind.cv)
    atoms = extract_atoms(first)
    assert first.document_id == second.document_id
    assert [atom.evidence_id for atom in atoms] == [atom.evidence_id for atom in extract_atoms(second)]
    assert grounding_problems(profile_for(first, atoms)) == []
    role = next(atom for atom in atoms if atom.employer == "Acme")
    assert role.title == "Engineer"
    assert role.start_date is not None and role.start_date.year == 2020
    bullet = next(atom for atom in atoms if "20%" in atom.statement)
    assert bullet.ownership_verb == "Led"
    assert bullet.metrics == ["20%"]
    assert "Postgres" in bullet.skills
    assert "PostgreSQL" not in bullet.skills


def test_multi_column_docx_is_refused(tmp_path: Path):
    path = tmp_path / "columns.docx"
    _write_docx(path, columns=2)
    with pytest.raises(UnsupportedLayoutError):
        parse_document(path, DocumentKind.cv)


def test_single_column_pdf_parses(tmp_path: Path):
    path = tmp_path / "cv.pdf"
    _write_pdf(
        path,
        [
            "Ada Lovelace",
            "PROFESSIONAL EXPERIENCE",
            "Engineer | Acme 2020 - 2022",
            "- Led the API migration and cut latency by 20% using Postgres.",
            "TECHNICAL SKILLS",
            "- Python, Postgres, Docker",
        ],
    )
    record = parse_document(path, DocumentKind.cv)
    atoms = extract_atoms(record)
    assert grounding_problems(profile_for(record, atoms)) == []
    assert any(atom.employer == "Acme" for atom in atoms)
    assert any(atom.metrics == ["20%"] for atom in atoms)


def test_multi_column_pdf_is_refused(tmp_path: Path):
    path = tmp_path / "columns.pdf"
    _write_two_column_pdf(path)
    with pytest.raises(UnsupportedLayoutError):
        parse_document(path, DocumentKind.cv)


def test_pdf_without_a_text_layer_is_refused(tmp_path: Path):
    path = tmp_path / "scan.pdf"
    blank = canvas.Canvas(str(path), pagesize=A4)
    blank.showPage()
    blank.save()
    with pytest.raises(ParseError, match="text layer"):
        parse_document(path, DocumentKind.cv)


def test_unsupported_extension_is_refused(tmp_path: Path):
    path = tmp_path / "notes.txt"
    path.write_text("hello", encoding="utf-8")
    with pytest.raises(ParseError, match="Supported files"):
        parse_document(path, DocumentKind.cv)
